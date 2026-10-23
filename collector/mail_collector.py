"""
Mail Feed collector - Kerem Demirbas
Reads new Outlook mails (READ-ONLY) and writes them as plain text.
It never sends, replies, moves, deletes or marks anything in Outlook.
Afterwards it pushes the text feed to GitHub, where analyzer/ runs on it.

Output (repo root = folder above this script):
  feed/YYYY-MM-DD/index.txt        one short line per mail      (pushed)
  feed/YYYY-MM-DD/mails/0001.txt   clean mail text + attachment text (pushed)
  feed/YYYY-MM-DD/files/...        original attachments          (local only)
  collector/state.json             what was already collected    (local only)
  collector/run_log.txt            one line per run              (local only)
"""
import os, re, sys, json, hashlib, subprocess, urllib.request, urllib.error
import datetime as dt, traceback, zipfile

# ---------------- settings ----------------
FOLDERS = ["Inbox", "Sent"]      # Inbox includes its subfolders
FIRST_RUN_DAYS = 3               # how far back the very first run looks
OVERLAP_DAYS = 2                 # re-check window (already-seen mails are skipped)
MAX_BODY = 15000                 # chars per mail body
MAX_ATT = 12000                  # chars of text per attachment
SKIP_SMALL_IMAGES = 30 * 1024    # logos / signatures
PUSH_TO_GITHUB = True            # git push feed/ after collecting (private repo only)
# -----------------------------------------

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
STATE = os.path.join(BASE, "state.json")
LOG = os.path.join(BASE, "run_log.txt")
IMG = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".heic", ".webp"}

QUOTE_MARK = re.compile(
    r"^\s*(From|Kimden|Gönderen|Gönderildi|De|От)\s*:|^\s*-{3,}\s*(Original Message|Özgün İleti|Orijinal İleti)"
    r"|^\s*_{10,}\s*$|^On .{5,80} wrote:\s*$|^.{5,80} tarihinde .{0,80} yazdı:\s*$",
    re.I | re.M)


def log(msg):
    with open(LOG, "a", encoding="utf-8") as f:
        f.write(f"{dt.datetime.now():%Y-%m-%d %H:%M}  {msg}\n")


def load_state():
    if os.path.exists(STATE):
        with open(STATE, encoding="utf-8") as f:
            return json.load(f)
    return {"last_run": None, "seen": {}, "conv": {}, "att": {}}


def save_state(s):
    cut = (dt.datetime.now() - dt.timedelta(days=60)).strftime("%Y-%m-%d")
    for k in ("seen", "conv", "att"):
        s[k] = {a: b for a, b in s[k].items() if b >= cut}
    tmp = STATE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(s, f)
    os.replace(tmp, STATE)


def pydt(t):
    return dt.datetime(t.year, t.month, t.day, t.hour, t.minute, t.second)


def clean(text):
    text = (text or "").replace("\r", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def safe(name):
    return re.sub(r'[\\/:*?"<>|]+', "_", name)[:120]


def sender(item):
    try:
        if item.SenderEmailType == "EX":
            u = item.Sender.GetExchangeUser()
            if u:
                return u.PrimarySmtpAddress
    except Exception:
        pass
    try:
        return item.SenderEmailAddress or ""
    except Exception:
        return ""


# ---------- attachment -> text ----------
def pdf_text(p):
    from pypdf import PdfReader
    r = PdfReader(p)
    return "\n".join((pg.extract_text() or "") for pg in r.pages[:20])


def xlsx_text(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 400:
                out.append("... more rows")
                break
            vals = [str(v).strip() for v in row if v not in (None, "")]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


def xls_text(p):
    import xlrd
    wb = xlrd.open_workbook(p)
    out = []
    for sh in wb.sheets():
        out.append(f"[sheet {sh.name}]")
        for r in range(min(sh.nrows, 400)):
            vals = [str(v).strip() for v in sh.row_values(r) if str(v).strip()]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


def docx_text(p):
    import docx
    d = docx.Document(p)
    out = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def msg_text(p, ns):
    m = ns.OpenSharedItem(p)
    t = f"[attached mail] {m.Subject}\nFrom: {getattr(m, 'SenderName', '')}\n{clean(m.Body)}"
    try:
        m.Close(1)  # olDiscard - nothing saved
    except Exception:
        pass
    return t


def att_text(p, ns):
    ext = os.path.splitext(p)[1].lower()
    if ext == ".pdf":
        return pdf_text(p)
    if ext in (".xlsx", ".xlsm"):
        return xlsx_text(p)
    if ext == ".xls":
        return xls_text(p)
    if ext == ".docx":
        return docx_text(p)
    if ext in (".txt", ".csv"):
        return open(p, encoding="utf-8", errors="ignore").read()
    if ext == ".msg":
        return msg_text(p, ns)
    if ext == ".zip":
        return "zip contents: " + ", ".join(zipfile.ZipFile(p).namelist()[:50])
    if ext in IMG:
        return None
    return None


# ---------- main ----------
def walk(folder):
    yield folder
    try:
        for f in folder.Folders:
            yield from walk(f)
    except Exception:
        return


def main():
    import win32com.client
    ol = win32com.client.Dispatch("Outlook.Application")
    ns = ol.GetNamespace("MAPI")
    st = load_state()
    now = dt.datetime.now()
    if st["last_run"]:
        since = dt.datetime.fromisoformat(st["last_run"]) - dt.timedelta(days=OVERLAP_DAYS)
    else:
        since = now - dt.timedelta(days=FIRST_RUN_DAYS)

    roots = []
    if "Inbox" in FOLDERS:
        roots += [(f, "IN") for f in walk(ns.GetDefaultFolder(6))]
    if "Sent" in FOLDERS:
        roots.append((ns.GetDefaultFolder(5), "OUT"))

    found = []
    for folder, kind in roots:
        try:
            items = folder.Items
            items.Sort("[ReceivedTime]", True)
        except Exception:
            continue
        for it in items:
            try:
                if it.Class != 43:  # mail items only
                    continue
                t = pydt(it.ReceivedTime if kind == "IN" else it.SentOn)
            except Exception:
                continue
            if t < since:
                break
            try:
                eid = it.EntryID
            except Exception:
                continue
            if eid in st["seen"]:
                continue
            found.append((t, kind, folder.Name, it, eid))

    found.sort(key=lambda x: x[0])
    day = now.strftime("%Y-%m-%d")
    out = os.path.join(ROOT, "feed", day)
    mdir, fdir = os.path.join(out, "mails"), os.path.join(out, "files")
    os.makedirs(mdir, exist_ok=True)
    os.makedirs(fdir, exist_ok=True)
    n0 = len([x for x in os.listdir(mdir) if x.endswith(".txt")])
    idx_lines = []

    for i, (t, kind, fname, it, eid) in enumerate(found, start=n0 + 1):
        no = f"{i:04d}"
        try:
            subj = it.Subject or ""
            frm = f"{it.SenderName} <{sender(it)}>"
            to, cc = it.To or "", it.CC or ""
            try:
                conv = it.ConversationID or subj
            except Exception:
                conv = subj
            body = clean(it.Body)
            note = ""
            if conv in st["conv"]:
                m = QUOTE_MARK.search(body)
                if m and m.start() > 0:
                    body = body[:m.start()].strip()
                    note = f"(earlier thread already collected on {st['conv'][conv]})"
            if len(body) > MAX_BODY:
                body = body[:MAX_BODY] + "\n... [cut]"

            att_names, att_blocks = [], []
            for a in it.Attachments:
                try:
                    an = a.FileName
                except Exception:
                    continue
                ext = os.path.splitext(an)[1].lower()
                path = os.path.join(fdir, f"{no}_{safe(an)}")
                try:
                    a.SaveAsFile(path)
                except Exception:
                    att_names.append(f"{an}(could not save)")
                    continue
                size = os.path.getsize(path)
                if ext in IMG and size < SKIP_SMALL_IMAGES:
                    os.remove(path)
                    continue
                h = hashlib.sha1(open(path, "rb").read()).hexdigest()
                if h in st["att"]:
                    att_names.append(f"{an}(seen before)")
                    att_blocks.append(f"--- attachment: {an} --- same file already collected on {st['att'][h]}")
                    os.remove(path)
                    continue
                st["att"][h] = day
                att_names.append(an)
                try:
                    txt = att_text(path, ns)
                except Exception as e:
                    txt = f"[could not read: {type(e).__name__}]"
                if txt is None:
                    txt = "[image or other file - saved, not read]"
                txt = clean(txt)
                if len(txt) > MAX_ATT:
                    txt = txt[:MAX_ATT] + "\n... [cut]"
                att_blocks.append(f"--- attachment: {an} (file: files/{os.path.basename(path)}) ---\n{txt}")

            with open(os.path.join(mdir, no + ".txt"), "w", encoding="utf-8") as f:
                f.write(f"No: {no}\nDate: {t:%d.%m.%Y %H:%M}\nFolder: {kind} / {fname}\n"
                        f"From: {frm}\nTo: {to}\nCC: {cc}\nSubject: {subj}\n{note}\n\n{body}\n\n"
                        + "\n\n".join(att_blocks))
            first = re.sub(r"\s+", " ", body)[:200]
            idx_lines.append(f"{no} | {t:%d.%m %H:%M} | {kind} | {frm} | to: {to[:80]} | {subj[:120]} | "
                             f"att: {', '.join(att_names) or '-'} | {first}")
            st["seen"][eid] = day
            st["conv"].setdefault(conv, day)
        except Exception as e:
            idx_lines.append(f"{no} | ERROR reading mail: {type(e).__name__}")

    if idx_lines:
        with open(os.path.join(out, "index.txt"), "a", encoding="utf-8") as f:
            f.write("\n".join(idx_lines) + "\n")
    st["last_run"] = now.isoformat(timespec="seconds")
    save_state(st)
    log(f"OK  {len(found)} new mails -> feed/{day}")
    if PUSH_TO_GITHUB:
        log(push_feed(day, len(found)))


# ---------- push to GitHub ----------
def git(*args):
    r = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, timeout=180,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if r.returncode != 0:
        raise RuntimeError(f"git {args[0]}: {(r.stderr or r.stdout).strip()[-300:]}")
    return r.stdout.strip()


def repo_is_private():
    """Unauthenticated API call: a public repo answers 200, a private one 404."""
    url = git("remote", "get-url", "origin")
    m = re.search(r"github\.com[:/]([^/]+)/([^/]+?)(\.git)?$", url)
    if not m:
        raise RuntimeError(f"origin is not a GitHub repo: {url}")
    try:
        urllib.request.urlopen(f"https://api.github.com/repos/{m[1]}/{m[2]}", timeout=20)
        return False
    except urllib.error.HTTPError as e:
        if e.code == 404:
            return True
        raise


def push_feed(day, n):
    try:
        if not repo_is_private():
            return "PUSH SKIPPED  repo is PUBLIC - make it private first (mails would be visible to everyone)"
        git("add", "feed")
        if git("diff", "--cached", "--name-only"):
            git("commit", "-m", f"feed: {n} new mails ({day})")
        # also retries an earlier commit whose push failed (offline etc.)
        git("pull", "--rebase", "--autostash", "-q")
        git("push", "-q")
        return "PUSH  OK"
    except Exception as e:
        return f"PUSH FAILED  {e}"


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        log("FAILED  " + repr(e) + " | " + traceback.format_exc().splitlines()[-1])
        sys.exit(1)
